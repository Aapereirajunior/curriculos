import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

from auto_jobs.models import TailoredApplication


def _safe_filename(text: str) -> str:
    return re.sub(r"[^a-zA-Z0-9_-]+", "_", text).strip("_")[:80]


def render_resume_docx(base_resume: dict, tailored: TailoredApplication, out_dir: str) -> str:
    doc = Document()

    doc.add_heading(base_resume.get("full_name", ""), level=0)
    doc.add_paragraph(base_resume.get("title", ""))

    contact = base_resume.get("contact", {})
    contact_line = " | ".join(
        v for v in [contact.get("email"), contact.get("phone"), contact.get("location")] if v
    )
    if contact_line:
        doc.add_paragraph(contact_line)

    links_line = " | ".join(
        v
        for v in [
            contact.get("linkedin"),
            contact.get("github"),
            contact.get("lattes"),
            contact.get("orcid"),
        ]
        if v
    )
    if links_line:
        doc.add_paragraph(links_line)

    doc.add_heading("Resumo", level=1)
    doc.add_paragraph(tailored.summary or base_resume.get("summary", ""))

    skills = base_resume.get("skills", [])
    if skills:
        doc.add_heading("Habilidades", level=1)
        doc.add_paragraph(", ".join(skills))

    doc.add_heading("Experiência", level=1)
    for exp in base_resume.get("experience", []):
        company = exp.get("company", "")
        p = doc.add_paragraph()
        run = p.add_run(f"{exp.get('role', '')} — {company}")
        run.bold = True
        run.font.size = Pt(11)
        doc.add_paragraph(f"{exp.get('start', '')} - {exp.get('end', '')}")

        bullets = tailored.tailored_bullets.get(company) or exp.get("bullets", [])
        for bullet in bullets:
            doc.add_paragraph(bullet, style="List Bullet")

    education = base_resume.get("education", [])
    if education:
        doc.add_heading("Educação", level=1)
        for edu in education:
            doc.add_paragraph(
                f"{edu.get('degree', '')} — {edu.get('institution', '')} "
                f"({edu.get('start', '')} - {edu.get('end', '')})"
            )

    certifications = base_resume.get("certifications", [])
    if certifications:
        doc.add_heading("Certificações", level=1)
        for cert in certifications:
            doc.add_paragraph(cert, style="List Bullet")

    publications = base_resume.get("publications", [])
    if publications:
        doc.add_heading("Produção Científica", level=1)
        for pub in publications:
            doc.add_paragraph(pub, style="List Bullet")

    presentations = base_resume.get("presentations", [])
    if presentations:
        doc.add_heading("Apresentações e Congressos", level=1)
        for pres in presentations:
            doc.add_paragraph(pres, style="List Bullet")

    languages = base_resume.get("languages", [])
    if languages:
        doc.add_heading("Idiomas", level=1)
        doc.add_paragraph(", ".join(languages))

    Path(out_dir).mkdir(parents=True, exist_ok=True)
    filename = f"resume_{_safe_filename(tailored.job.company)}_{_safe_filename(tailored.job.title)}.docx"
    out_path = str(Path(out_dir) / filename)
    doc.save(out_path)
    return out_path


def render_cover_letter_docx(tailored: TailoredApplication, out_dir: str) -> str:
    doc = Document()
    for paragraph in tailored.cover_letter.split("\n\n"):
        doc.add_paragraph(paragraph.strip())

    Path(out_dir).mkdir(parents=True, exist_ok=True)
    filename = f"cover_{_safe_filename(tailored.job.company)}_{_safe_filename(tailored.job.title)}.docx"
    out_path = str(Path(out_dir) / filename)
    doc.save(out_path)
    return out_path
